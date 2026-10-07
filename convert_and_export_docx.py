import json
import re
from opencc import OpenCC
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

cc = OpenCC('s2t')

def clean_brackets(text):
    # 移除題目末尾或中間的括號答案空格：（　　）、（ ）、()、（）、( ) 等
    # 處理全形與半形括號
    # 1. 結尾處的括弧
    text = re.sub(r'[（(]\s*[)）]\s*[。.]?\s*$', '', text)
    text = re.sub(r'[（(]\s*[\u3000\s]*[)）]\s*[。.]?\s*$', '', text)
    # 2. 句尾的（　　）
    text = re.sub(r'[（(]\s*[\u3000\s]*[)）]', '', text)
    # 修剪結尾多餘標點或空白
    text = text.rstrip()
    if not text.endswith(('。', '？', '！', '：', ':', '?', '!')):
        text += '。'
    return text

print("Converting questions.json to Traditional Chinese and removing parentheses...")
raw_questions = json.load(open('questions.json', 'r', encoding='utf-8'))

converted_questions = []
for q in raw_questions:
    new_q = dict(q)
    # 題幹繁體化並去除括號
    new_q['text'] = clean_brackets(cc.convert(q['text']))
    
    # 選項繁體化
    if 'options' in q:
        new_opts = []
        for opt in q['options']:
            new_opt = dict(opt)
            new_opt['text'] = cc.convert(opt['text'])
            new_opts.append(new_opt)
        new_q['options'] = new_opts
        
    # 簡答題答案繁體化
    if q['type'] == 'short' and isinstance(q.get('answer'), str):
        new_q['answer'] = cc.convert(q['answer'])
        
    converted_questions.append(new_q)

# 寫回 questions.json
with open('questions.json', 'w', encoding='utf-8') as f:
    json.dump(converted_questions, f, ensure_ascii=False, indent=1)

print("questions.json successfully updated!")

# 產生 docx 完整題庫核對檔（包含題幹、選項、標準答案）
print("Generating 完整題庫核對.docx...")
doc = docx.Document()

# 頁面邊界設定
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# 大標題
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title_p.add_run("羽絨行業職業技能理論考試題庫及參考答案（完整核對版）")
run.font.size = Pt(18)
run.font.bold = True
run.font.name = "微軟正黑體"

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_run = sub_p.add_run("收錄題型：單選題（500題）、多選題（240題）、判斷題（230題）、簡答題（30題）｜共1000題")
sub_run.font.size = Pt(11)
sub_run.font.color.rgb = RGBColor(100, 100, 100)
sub_run.font.name = "微軟正黑體"

type_titles = {
    'single': '一、單選題（共 500 題）',
    'multi': '二、多選題（共 240 題）',
    'judge': '三、判斷題（共 230 題）',
    'short': '四、簡答題（共 30 題）'
}

by_type = {'single': [], 'multi': [], 'judge': [], 'short': []}
for q in converted_questions:
    by_type[q['type']].append(q)

for qtype in ['single', 'multi', 'judge', 'short']:
    doc.add_paragraph() # 空行
    h = doc.add_paragraph()
    h_run = h.add_run(type_titles[qtype])
    h_run.font.size = Pt(14)
    h_run.font.bold = True
    h_run.font.color.rgb = RGBColor(40, 70, 120)
    h_run.font.name = "微軟正黑體"
    
    for q in by_type[qtype]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        
        # 題號與題幹
        num_run = p.add_run(f"{q['no']}. ")
        num_run.bold = True
        num_run.font.name = "微軟正黑體"
        
        text_run = p.add_run(q['text'])
        text_run.font.name = "微軟正黑體"
        
        # 答案呈現
        if qtype == 'single':
            ans_str = f"  【標準答案：{q['answer']}】"
            ans_run = p.add_run(ans_str)
            ans_run.bold = True
            ans_run.font.color.rgb = RGBColor(30, 130, 60)
            ans_run.font.name = "微軟正黑體"
            
            # 選項列表
            for opt in q['options']:
                opt_p = doc.add_paragraph()
                opt_p.paragraph_format.left_indent = Inches(0.25)
                opt_p.paragraph_format.space_after = Pt(1)
                opt_run = opt_p.add_run(f"{opt['key']}. {opt['text']}")
                opt_run.font.name = "微軟正黑體"
                if opt['key'] == q['answer']:
                    opt_run.font.bold = True
                    opt_run.font.color.rgb = RGBColor(30, 130, 60)
                    
        elif qtype == 'multi':
            ans_join = "".join(q['answer'])
            ans_str = f"  【標準答案：{ans_join}】"
            ans_run = p.add_run(ans_str)
            ans_run.bold = True
            ans_run.font.color.rgb = RGBColor(30, 130, 60)
            ans_run.font.name = "微軟正黑體"
            
            for opt in q['options']:
                opt_p = doc.add_paragraph()
                opt_p.paragraph_format.left_indent = Inches(0.25)
                opt_p.paragraph_format.space_after = Pt(1)
                opt_run = opt_p.add_run(f"{opt['key']}. {opt['text']}")
                opt_run.font.name = "微軟正黑體"
                if opt['key'] in q['answer']:
                    opt_run.font.bold = True
                    opt_run.font.color.rgb = RGBColor(30, 130, 60)
                    
        elif qtype == 'judge':
            ans_display = "正確 (√)" if q['answer'] == 'T' else "錯誤 (×)"
            ans_str = f"  【標準答案：{ans_display}】"
            ans_run = p.add_run(ans_str)
            ans_run.bold = True
            ans_run.font.color.rgb = RGBColor(30, 130, 60)
            ans_run.font.name = "微軟正黑體"
            
        elif qtype == 'short':
            ans_p = doc.add_paragraph()
            ans_p.paragraph_format.left_indent = Inches(0.25)
            ans_p.paragraph_format.space_after = Pt(4)
            ref_head = ans_p.add_run("【參考答案】：\n")
            ref_head.bold = True
            ref_head.font.color.rgb = RGBColor(30, 130, 60)
            ref_head.font.name = "微軟正黑體"
            
            ans_body = ans_p.add_run(q['answer'])
            ans_body.font.color.rgb = RGBColor(60, 60, 60)
            ans_body.font.name = "微軟正黑體"

doc_path = "KLF_羽絨行業理論考試題庫_完整1000題.docx"
doc.save(doc_path)
print(f"Docx file saved successfully as {doc_path}!")
