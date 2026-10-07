import json
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

print("Generating docx in Simplified Chinese...")
questions = json.load(open('questions.json', 'r', encoding='utf-8'))

doc = docx.Document()

# 页面边距设置
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# 大标题
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title_p.add_run("羽绒行业职业技能理论考试题库及参考答案（完整核对版）")
run.font.size = Pt(18)
run.font.bold = True
run.font.name = "微软雅黑"

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_run = sub_p.add_run("收录题型：单选题（500题）、多选题（240题）、判断题（230题）、简答题（30题）｜ 共1000题")
sub_run.font.size = Pt(11)
sub_run.font.color.rgb = RGBColor(100, 100, 100)
sub_run.font.name = "微软雅黑"

type_titles = {
    'single': '一、单选题（共 500 题）',
    'multi': '二、多选题（共 240 题）',
    'judge': '三、判断题（共 230 题）',
    'short': '四、简答题（共 30 题）'
}

by_type = {'single': [], 'multi': [], 'judge': [], 'short': []}
for q in questions:
    by_type[q['type']].append(q)

for qtype in ['single', 'multi', 'judge', 'short']:
    doc.add_paragraph() # 空行
    h = doc.add_paragraph()
    h_run = h.add_run(type_titles[qtype])
    h_run.font.size = Pt(14)
    h_run.font.bold = True
    h_run.font.color.rgb = RGBColor(40, 70, 120)
    h_run.font.name = "微软雅黑"
    
    for q in by_type[qtype]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        
        # 题号与题干
        num_run = p.add_run(f"{q['no']}. ")
        num_run.bold = True
        num_run.font.name = "微软雅黑"
        
        text_run = p.add_run(q['text'])
        text_run.font.name = "微软雅黑"
        
        # 答案呈现
        if qtype == 'single':
            ans_str = f"  【标准答案：{q['answer']}】"
            ans_run = p.add_run(ans_str)
            ans_run.bold = True
            ans_run.font.color.rgb = RGBColor(30, 130, 60)
            ans_run.font.name = "微软雅黑"
            
            # 选项列表
            for opt in q['options']:
                opt_p = doc.add_paragraph()
                opt_p.paragraph_format.left_indent = Inches(0.25)
                opt_p.paragraph_format.space_after = Pt(1)
                opt_run = opt_p.add_run(f"{opt['key']}. {opt['text']}")
                opt_run.font.name = "微软雅黑"
                if opt['key'] == q['answer']:
                    opt_run.font.bold = True
                    opt_run.font.color.rgb = RGBColor(30, 130, 60)
                    
        elif qtype == 'multi':
            ans_join = "".join(q['answer'])
            ans_str = f"  【标准答案：{ans_join}】"
            ans_run = p.add_run(ans_str)
            ans_run.bold = True
            ans_run.font.color.rgb = RGBColor(30, 130, 60)
            ans_run.font.name = "微软雅黑"
            
            for opt in q['options']:
                opt_p = doc.add_paragraph()
                opt_p.paragraph_format.left_indent = Inches(0.25)
                opt_p.paragraph_format.space_after = Pt(1)
                opt_run = opt_p.add_run(f"{opt['key']}. {opt['text']}")
                opt_run.font.name = "微软雅黑"
                if opt['key'] in q['answer']:
                    opt_run.font.bold = True
                    opt_run.font.color.rgb = RGBColor(30, 130, 60)
                    
        elif qtype == 'judge':
            ans_display = "正确 (√)" if q['answer'] == 'T' else "错误 (×)"
            ans_str = f"  【标准答案：{ans_display}】"
            ans_run = p.add_run(ans_str)
            ans_run.bold = True
            ans_run.font.color.rgb = RGBColor(30, 130, 60)
            ans_run.font.name = "微软雅黑"
            
        elif qtype == 'short':
            ans_p = doc.add_paragraph()
            ans_p.paragraph_format.left_indent = Inches(0.25)
            ans_p.paragraph_format.space_after = Pt(4)
            ref_head = ans_p.add_run("【参考答案】：\n")
            ref_head.bold = True
            ref_head.font.color.rgb = RGBColor(30, 130, 60)
            ref_head.font.name = "微软雅黑"
            
            ans_body = ans_p.add_run(q['answer'])
            ans_body.font.color.rgb = RGBColor(60, 60, 60)
            ans_body.font.name = "微软雅黑"

doc_path = "KLF_羽绒行业理论考试题库_完整1000题.docx"
doc.save(doc_path)
print("Docx export in Simplified Chinese completed successfully!")
